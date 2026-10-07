#!/usr/bin/env python3
"""Generate Docs/KAERI_Modelling_Direction_2026.pptx (Appendix J-style report slides)."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

REPO_ROOT = Path(__file__).resolve().parents[1]
OUT_PATH = REPO_ROOT / "Docs" / "KAERI_Modelling_Direction_2026.pptx"

UG = REPO_ROOT / "Docs" / "User's Guide" / "Imagery"
KURT = Path(
    r"C:\_repo\Work_v2\KURT_Model\2_models\1_Regional_Boundary_Average_Rainfall\Docs\imagery"
)
KURT_RTS = Path(
    r"C:\_repo\Work_v2\KURT_Model\2_models\2_TestTransientRainfall\_TecplotReview"
)
KURT_GSTR = Path(
    r"C:\_repo\Work_v2\KURT_Model\2_models\2_TestTransientRainfall\4_GSTR\Docs_meeting"
)
EX_ABDUL = Path(r"C:\_repo\Grdbldr\MUT_Examples\Verification\6_Abdul_Prism_Cell")
EX_CLN = Path(r"C:\_repo\Grdbldr\MUT_Examples\Verification\3_1_CLN_for_SWF")

# 4:3 report size (Appendix J)
SLIDE_W = Inches(10.0)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.25)
TITLE_TOP = Inches(0.18)
TITLE_H = Inches(0.55)
BODY_TOP = Inches(0.85)
FOOTER_Y = Inches(7.15)
ACCENT = RGBColor(0x1B, 0x3A, 0x5F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BODY = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x55, 0x55, 0x55)
CAPTION = RGBColor(0x33, 0x33, 0x33)
DATE_STR = "11 September 2026"


def _font(run, *, size_pt: float, bold: bool = False, color: RGBColor = BODY, name: str = "Calibri"):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color


def _textbox(slide, left, top, width, height, text: str, *, size=14, bold=False, color=BODY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", "t" if anchor == MSO_ANCHOR.TOP else "ctr")
    except Exception:
        pass
    # Split paragraphs on blank lines
    parts = [p.strip() for p in text.split("\n\n") if p.strip()]
    if not parts:
        parts = [""]
    for i, part in enumerate(parts):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(8)
        # Allow single newlines within a paragraph as soft breaks via separate runs/lines
        lines = part.split("\n")
        for j, line in enumerate(lines):
            if j == 0:
                run = p.add_run()
                run.text = line
                _font(run, size_pt=size, bold=bold, color=color)
            else:
                # new paragraph for numbered items / list lines
                p2 = tf.add_paragraph()
                p2.alignment = align
                p2.space_before = Pt(4)
                p2.space_after = Pt(4)
                run = p2.add_run()
                run.text = line
                _font(run, size_pt=size, bold=bold, color=color)
    return box


def _title_bar(slide, title: str):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_W, TITLE_H)
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    _textbox(
        slide,
        MARGIN,
        Inches(0.08),
        SLIDE_W - 2 * MARGIN,
        Inches(0.45),
        title,
        size=18,
        bold=True,
        color=WHITE,
    )


def _footer(slide, page: int):
    _textbox(slide, MARGIN, FOOTER_Y, Inches(3.5), Inches(0.25), DATE_STR, size=9, color=MUTED)
    _textbox(
        slide,
        Inches(4.2),
        FOOTER_Y,
        Inches(1.5),
        Inches(0.25),
        str(page),
        size=9,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    _textbox(
        slide,
        Inches(6.0),
        FOOTER_Y,
        Inches(3.7),
        Inches(0.25),
        "KAERI modelling direction · MUT 2025.026",
        size=9,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )


def _fit_picture(slide, path: Path, left, top, max_w, max_h):
    """Place picture preserving aspect ratio inside the max box; return (pic, used_w, used_h)."""
    if not path.exists():
        raise FileNotFoundError(path)
    from PIL import Image

    with Image.open(path) as im:
        pw, ph = im.size
    aspect = pw / ph
    box_aspect = max_w / max_h
    if aspect >= box_aspect:
        used_w = max_w
        used_h = max_w / aspect
    else:
        used_h = max_h
        used_w = max_h * aspect
    pic = slide.shapes.add_picture(str(path), left, top, width=used_w, height=used_h)
    return pic, used_w, used_h


def _picture_with_caption(slide, path: Path, left, top, max_w, max_h, caption: str, cap_size=10):
    pic, used_w, used_h = _fit_picture(slide, path, left, top, max_w, max_h)
    cap_top = top + used_h + Inches(0.05)
    _textbox(slide, left, cap_top, max(used_w, max_w * 0.9), Inches(0.45), caption, size=cap_size, color=CAPTION)
    return pic


def _require(path: Path) -> Path:
    if not path.exists():
        raise FileNotFoundError(f"Missing figure: {path}")
    return path


# ---------------------------------------------------------------------------
# Slide content
# ---------------------------------------------------------------------------

SLIDES: list[dict] = [
    {
        "kind": "title",
        "title": "KAERI Modelling Approach",
        "subtitle": "Direction and 2026 development summary",
        "lines": [
            "Integrated groundwater, surface-water and conduit flow for KURT",
            "MUT · MODFLOW-USG · QGIS · Tecplot",
            "Current MUT release: 2025.026",
        ],
    },
    {
        "kind": "left_text_right_fig",
        "title": "Purpose - Supporting KURT engineering assessment",
        "prose": (
            "The KAERI/KURT assessment needs a reproducible path from site GIS into a coupled "
            "flow model that can represent regional groundwater, overland response to rainfall, "
            "and engineered conduits or shafts.\n\n"
            "QGIS organises DEMs and spatial layers; MUT translates those into MODFLOW-USG inputs; "
            "usgs_1 solves the flow system; Tecplot and the model dossier communicate water levels, "
            "exchange fluxes and budgets for engineering review."
        ),
        "image": UG / "QGIS_ProjectDEM.png",
        "caption": "Example project DEM prepared in QGIS for regional model construction.",
    },
    {
        "kind": "left_text_right_fig",
        "title": "Coupled domains - Groundwater, surface water and conduits",
        "prose": (
            "Engineering questions at KURT are rarely single-domain. Rainfall first wets the surface "
            "(SWF), infiltrates to the porous medium (GWF), and may interact with discrete conduits "
            "or shafts (CLN).\n\n"
            "The 2026 toolchain builds these domains on a shared mesh so exchange fluxes, heads and "
            "budgets can be examined together rather than as separate models stitched after the fact."
        ),
        "image": UG / "3_26_SWF_GWF.png",
        "caption": "SWF surface cells overlying the GWF mesh - dual-domain exchange geometry.",
    },
    {
        "kind": "left_text_right_fig",
        "title": "Rainfall and surface forcing - Zones, transient recharge and controls",
        "prose": (
            "Surface hydrology for engineering design needs spatially varying rainfall and controllable "
            "boundary heads. MUT now supports multi-zone SWF transient recharge (merged RTS columns and "
            "matching RCH zone maps), reports whether each stress period uses RCH or RTS, and writes "
            "CLN/SWF constant-head boundaries as start-to-end head ramps so engineered water levels can "
            "be imposed consistently through a simulation."
        ),
        "image": UG / "3_35_SWFRecharge.png",
        "caption": "Applied SWF recharge mapped on the surface domain for review in Tecplot.",
    },
    {
        "kind": "left_text_right_fig",
        "title": "Spatio-temporal recharge - Raster snapshots for transient rain",
        "prose": (
            "For storms and spatially migrating rainfall fields, General Spatio-Temporal Recharge (GSTR) "
            "assigns ArcASCII raster snapshots to named GWF, SWF or CLN instances and interpolates rates "
            "in time. Existing RCH/RTS packages remain available; when both are active the fluxes add.\n\n"
            "A KURT demonstration lives under the transient-rainfall GSTR test model. The figure shows "
            "applied recharge from the related RTS workflow used to evaluate surface response."
        ),
        "image": KURT_RTS / "1_RTS_Applied_RCH_with_legend.png",
        "caption": "Applied recharge field from the KURT transient-rainfall RTS review (proxy for GSTR-style spatial forcing).",
    },
    {
        "kind": "top_text_two_figs",
        "title": "KURT 4_GSTR example - Circular storm applied to SWF",
        "prose": (
            "The TestTransientRainfall/4_GSTR model drives SWF recharge with named GSTR instance StormEW: "
            "ArcASCII storm rasters are interpolated in time and applied on the surface domain. Post-processing "
            "writes cell-by-cell applied rates (_posto.modflow.SWF_GSTR.tecplot.dat) and a domain-mean time series. "
            "The maps below show the peak applied-rate field (t = 0.375 d) and the short storm pulse in mean GSTR "
            "before rates return to zero for the remainder of the run."
        ),
        "images": [
            (KURT_GSTR / "GSTR_AppliedRate_Map.png", "Applied SWF GSTR rate at peak mean (StormEW, t = 0.375 d)"),
            (KURT_GSTR / "GSTR_MeanRate_vs_Time.png", "Domain-mean applied GSTR rate vs time (first 2 days)"),
        ],
    },
    {
        "kind": "left_text_right_fig",
        "title": "Evapotranspiration in MODFLOW-USG - Head-dependent plant and soil water loss",
        "prose": (
            "usgs_1 (MODFLOW-USG / USG-Transport) already provides classical groundwater ET packages that "
            "remove water from the saturated zone as a function of water-table position. The EVT package "
            "applies a maximum ET rate that declines linearly to zero over an extinction depth below a "
            "specified ET surface. The ETS package extends this with a segmented (piecewise) ET-versus-depth "
            "curve for non-linear plant uptake.\n\n"
            "EVT can also take time-varying maximum rates (ETS keyword / time-series option) so seasonal "
            "potential ET can be imposed. These packages act on GWF nodes and appear as budget terms when "
            "activated - the practical path today for regional ET sinks alongside rainfall recharge."
        ),
        "image": KURT / "GWF_WaterTable.png",
        "caption": "Regional water-table surface - ET demand in EVT/ETS is controlled by proximity to the land surface / extinction depth.",
    },
    {
        "kind": "left_text_right_fig",
        "title": "Evapotranspiration in MUT - Database scaffolding, writers still to come",
        "prose": (
            "MUT can load an ET parameter database (et database) whose fields mirror an HGS-style plant/soil "
            "ET concept: evaporation and root depths, LAI tables, wilting point and field capacity, oxic/anoxic "
            "limits, and canopy interception storage. Sparse ET material IDs are supported alongside GWF/CLN/SWF "
            "databases.\n\n"
            "However, ET.xlsx / LAI.xlsx remain placeholders for future development: MUT does not yet assign ET "
            "zones to cells or write EVT/ETS package files. Name-file recognition of EVT/EVS/ETS exists for reading "
            "external datasets, but the KURT workflow today is rainfall and infiltration without an automated "
            "MUT-driven ET sink. Closing the gap means mapping the ET database (or a simpler EVT/ETS instruction "
            "set) onto usgs_1 package writers and exposing ET in volume-budget and Tecplot review."
        ),
        "image": KURT / "GWF_VolumeBudget.png",
        "caption": "Current KURT volume-budget dossier frame - ET package terms appear only when EVT/ETS are active in usgs_1.",
    },
    {
        "kind": "left_text_right_fig",
        "title": "Conduits and shafts - Building CLN from the mesh",
        "prose": (
            "Engineered openings and discrete pathways are represented with the Connected Linear Network "
            "(CLN). Development this year focused on intersecting conduit polylines with GWF/SWF element "
            "faces, preserving network connectivity, and connecting CLN nodes to the groundwater mesh for "
            "simple and shaft-style configurations.\n\n"
            "Conduits can also be rebuilt from a structure file with a companion coordinate list for geometry checks."
        ),
        "image": UG / "3_36_CLN.png",
        "caption": "CLN conduits shown with the surrounding porous-medium mesh.",
    },
    {
        "kind": "left_text_right_fig",
        "title": "CLN hydraulic detailing - Skin, geometry and initial heads",
        "prose": (
            "Once the conduit geometry exists, hydraulic behaviour must match the engineering concept. "
            "Per-cell skin conductivity (FSKIN) defaults from connected GWF horizontal conductivity, "
            "general tabular cross-sections support non-circular conduits, and initial CLN heads can be "
            "set from the land surface via column tops - a practical starting condition for shafts that "
            "daylight at topography."
        ),
        "image": UG / "MUT_InitialCLNHead.png",
        "caption": "CLN initial heads assigned from surface elevation / GWF column tops.",
    },
    {
        "kind": "left_text_right_fig",
        "title": "Velocity and pathways - Darcy and average-linear flow",
        "prose": (
            "Flow interpretation and future solute transport both need velocities, not only heads. Builds "
            "now emit a default VEL package for active GWF, CLN and SWF domains. Post-processing writes "
            "Darcy and average-linear velocity components for Tecplot; porosity (and CLN infill porosity "
            "where used) converts Darcy flux to pore velocity.\n\n"
            "An optional original SWF Manning velocity path remains available for continuity with HGS-style surface flow."
        ),
        "image": KURT / "GWF_WaterTable.png",
        "caption": "Regional water-table iso-surface from the KURT average-rainfall model dossier.",
    },
    {
        "kind": "left_text_right_fig",
        "title": "Communicating results - Tecplot SZL and model dossiers",
        "prose": (
            "Engineering review depends on clear maps of saturation, infiltration, water table and volume "
            "budgets. Finite-element results default to compact TecIO SZL files with all output times.\n\n"
            "After build or post-processing, mut_document writes a model dossier PDF and sectioned Tecplot "
            "layouts - including dedicated views for volume budget, water table, SWF-to-GWF infiltration "
            "and saturation slices - so reviewers see the same figures the modeller used."
        ),
        "image": UG / "4_16_Infiltration.png",
        "caption": "Example SWF-to-GWF infiltration map used in post-processed hillslope review.",
    },
    {
        "kind": "top_text_three_figs",
        "title": "Regional KURT illustration - Mesh, water table and infiltration",
        "prose": (
            "The regional average-rainfall KURT model exercises the full stack on site-scale geometry: "
            "dual-domain mesh, transient surface response, and groundwater storage. Side-by-side dossier "
            "frames below illustrate how mesh structure, water-table position and infiltration patterns "
            "are delivered together for engineering discussion."
        ),
        "images": [
            (KURT / "GWF_Mesh.png", "GWF mesh / domain overview"),
            (KURT / "GWF_WaterTable.png", "Water-table surface"),
            (KURT / "SWF_Infiltration.png", "SWF-to-GWF infiltration"),
        ],
    },
    {
        "kind": "top_text_two_figs",
        "title": "Verification practice - Confidence from known problems",
        "prose": (
            "Before relying on regional forecasts, the same toolchain is checked against published and "
            "in-house flow problems (Abdul hillslope, SWF critical-depth and CLN-for-SWF cases). Outflow "
            "and depth comparisons against accepted solutions provide engineering confidence that "
            "surface–subsurface coupling and conduit representations behave as intended when rainfall, "
            "outlets and materials are known."
        ),
        "images": [
            (EX_ABDUL / "Outflow_Comparison.png", "Abdul outlet hydrograph comparison"),
            (EX_CLN / "Outflow_Comparison.png", "CLN-for-SWF outflow comparison"),
        ],
    },
    {
        "kind": "left_text_right_fig",
        "title": "Future capability - Discrete Fracture Network (DFN) medium",
        "prose": (
            "KURT fracture zones are naturally represented as discrete 2D flow surfaces embedded in "
            "3D porous rock, not as continuum smearing alone and not as 1D pipes. A new medium, DFN, "
            "would sit alongside GWF, CLN and SWF: planar fracture elements with aperture and "
            "transmissivity, exchanging water (and later solute) with the surrounding GWF matrix.\n\n"
            "CLN remains the right tool for shafts and tubular conduits; DFN targets sheet-like "
            "fractures imported from site STL/GIS fracture surfaces (as in the KURT incoming fracture set)."
        ),
        "image": Path(r"C:\_repo\Work_v2\KURT_Model\0_incoming\Fracture_STL\fractures.png"),
        "caption": "KURT fracture surfaces (incoming STL set) - candidate geometry for a DFN medium.",
    },
    {
        "kind": "left_text_right_fig",
        "title": "Implementing DFN - Steps for 2D fracture flow and transport",
        "prose": (
            "1. Design - Equations, aperture/T properties, DFN-GWF exchange; NAM type and IDOMAIN code\n"
            "2. usgs_1 flow - Domain arrays, matrix assembly, budgets, binary heads/fluxes (patterned on CLN/SWF)\n"
            "3. Geometry - Import fracture surfaces; intersect/mesh onto GWF; write DFN GSF connectivity\n"
            "4. MUT - Domain, DFN materials DB, _build.mut verbs, package writers, Tecplot/dossier views\n"
            "5. Transport - BCT (or equivalent) on DFN nodes; matrix exchange; CON save and observations\n"
            "6. Verify - Single-fracture flow then transport; multi-fracture KURT case against a known solution"
        ),
        "image": UG / "3_38_CLN_GWF_Cells.png",
        "caption": "Analogue: CLN-GWF coupling today - DFN would add planar fracture-matrix exchange in the same multi-domain stack.",
    },
    {
        "kind": "left_text_right_fig",
        "title": "Gaps and next direction - Aligning solvers, ET writers and transport",
        "prose": (
            "The flow modelling stack is usable today at MUT 2025.026, but a few pieces still need aligned "
            "USG builds: velocity post-processing requires a usgs_1 that recognises the VEL name-file type, "
            "and GSTR rates must be solved by a USG-Beta build that includes the GSTR package. MUT ET package "
            "writers (EVT/ETS from the existing ET database or a simpler instruction set) are still future work.\n\n"
            "Further ahead: SWF/GWF solute transport (Abdul path) and a new DFN medium for 2D fracture flow "
            "and transport on KURT fracture surfaces."
        ),
        "image": UG / "4_5_Post_Abdul.png",
        "caption": "Abdul hillslope post-processed results - candidate platform for transport verification.",
    },
    {
        "kind": "left_text_right_fig",
        "title": "Discussion - Priorities for the next phase",
        "prose": (
            "With the 2026 flow capabilities in place, the meeting should set priorities for KURT delivery:\n"
            "1. Deepen rainfall and GSTR workflows for regional transient forcing\n"
            "2. Add MUT writers for MODFLOW EVT/ETS (or fuller HGS-style ET) so regional water balances include plant/soil water loss\n"
            "3. Extend shaft / CLN engineering detail on the regional mesh\n"
            "4. Advance surface-subsurface transport (Abdul verification path)\n"
            "5. Scope DFN (2D discrete fracture flow and transport) against KURT fracture geometry\n"
            "6. Agree which usgs_1 build is the project standard for VEL and GSTR"
        ),
        "image": KURT / "GWF_Results.png",
        "caption": "Example regional GWF results frame from the KURT model dossier.",
    },
]


def _render_title(slide, spec: dict, page: int):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = ACCENT
    bg.line.fill.background()
    strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(4.2), SLIDE_W, Inches(0.07))
    strip.fill.solid()
    strip.fill.fore_color.rgb = RGBColor(0x2E, 0x5A, 0x88)
    strip.line.fill.background()
    _textbox(slide, MARGIN, Inches(2.0), SLIDE_W - 2 * MARGIN, Inches(0.7), spec["title"], size=32, bold=True, color=WHITE)
    _textbox(slide, MARGIN, Inches(2.8), SLIDE_W - 2 * MARGIN, Inches(0.5), spec["subtitle"], size=18, color=RGBColor(0xD0, 0xDC, 0xE8))
    lines = "\n".join(spec["lines"])
    _textbox(slide, MARGIN, Inches(4.5), SLIDE_W - 2 * MARGIN, Inches(2.0), lines, size=16, color=WHITE)
    _textbox(slide, MARGIN, FOOTER_Y, Inches(4), Inches(0.25), DATE_STR, size=9, color=RGBColor(0xB0, 0xC0, 0xD0))
    _textbox(slide, Inches(8.5), FOOTER_Y, Inches(1.2), Inches(0.25), str(page), size=9, color=RGBColor(0xB0, 0xC0, 0xD0), align=PP_ALIGN.RIGHT)


def _render_left_right(slide, spec: dict, page: int):
    _title_bar(slide, spec["title"])
    text_w = Inches(3.6)
    _textbox(slide, MARGIN, BODY_TOP, text_w, Inches(5.9), spec["prose"], size=13)
    fig_left = Inches(4.0)
    fig_top = BODY_TOP
    fig_w = Inches(5.7)
    fig_h = Inches(5.5)
    _picture_with_caption(slide, _require(spec["image"]), fig_left, fig_top, fig_w, fig_h, spec["caption"])
    _footer(slide, page)


def _render_top_three(slide, spec: dict, page: int):
    _title_bar(slide, spec["title"])
    _textbox(slide, MARGIN, BODY_TOP, SLIDE_W - 2 * MARGIN, Inches(1.35), spec["prose"], size=13)
    fig_top = Inches(2.4)
    fig_h = Inches(4.0)
    gap = Inches(0.15)
    col_w = (SLIDE_W - 2 * MARGIN - 2 * gap) / 3
    left = MARGIN
    for path, caption in spec["images"]:
        _picture_with_caption(slide, _require(path), left, fig_top, col_w, fig_h, caption, cap_size=9)
        left = left + col_w + gap
    _footer(slide, page)


def _render_top_two(slide, spec: dict, page: int):
    _title_bar(slide, spec["title"])
    _textbox(slide, MARGIN, BODY_TOP, SLIDE_W - 2 * MARGIN, Inches(1.5), spec["prose"], size=13)
    fig_top = Inches(2.55)
    fig_h = Inches(4.0)
    gap = Inches(0.25)
    col_w = (SLIDE_W - 2 * MARGIN - gap) / 2
    left = MARGIN
    for path, caption in spec["images"]:
        _picture_with_caption(slide, _require(path), left, fig_top, col_w, fig_h, caption, cap_size=10)
        left = left + col_w + gap
    _footer(slide, page)


def build() -> Path:
    # Pillow is used for aspect-fit; fall back note if missing
    try:
        import PIL  # noqa: F401
    except ImportError as e:
        raise SystemExit("Pillow is required: pip install Pillow") from e

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    for i, spec in enumerate(SLIDES, start=1):
        slide = prs.slides.add_slide(blank)
        kind = spec["kind"]
        if kind == "title":
            _render_title(slide, spec, i)
        elif kind == "left_text_right_fig":
            _render_left_right(slide, spec, i)
        elif kind == "top_text_three_figs":
            _render_top_three(slide, spec, i)
        elif kind == "top_text_two_figs":
            _render_top_two(slide, spec, i)
        else:
            raise ValueError(kind)

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    try:
        prs.save(str(OUT_PATH))
        return OUT_PATH
    except PermissionError:
        alt = OUT_PATH.with_name(OUT_PATH.stem + "_new.pptx")
        prs.save(str(alt))
        print(f"Note: {OUT_PATH.name} is locked; wrote {alt.name} instead")
        return alt


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path} ({path.stat().st_size} bytes, {len(SLIDES)} slides)")
